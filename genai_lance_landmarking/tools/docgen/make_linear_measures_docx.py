"""Generates genai_lance_landmarking/production/Lance_v14_Linear_Measures.docx: definitions of the linear
measures computed by production/lance_v14_morphometrics.R (lateral_measures, bottom_measures), with
figures drawn on a real specimen.

Every figure is drawn from the CNN's v1.4 prediction for the example lance, using the same point numbers
and the same geometry as the R script (lance axis R02 -> R01, heights measured perpendicular to that axis
up to the dorsal curve, Bottom axis B11 -> apex midpoint), so the figures show exactly what the numbers
measure. Rerun after any change to the R script's measures rather than hand-editing the .docx.

usage (from repo root, project venv):
    python genai_lance_landmarking/tools/docgen/make_linear_measures_docx.py
"""
import csv
import sys
from pathlib import Path

import numpy as np
import tifffile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from measure_figs import Fig, PALETTE as C  # noqa: E402

ROOT = HERE.parents[1]
REPO = ROOT.parent
PROD = ROOT / "production"
OUT = PROD / "Lance_v14_Linear_Measures.docx"
FIGDIR = ROOT / "figures" / "linear_measures"
FIGDIR.mkdir(parents=True, exist_ok=True)
EXAMPLE = "PX014xNP060v4-2"  # N. pinetum parent, passes QC in all three views
KEY = list(csv.DictReader(open(PROD / "output" / "geomorph" / "landmark_key_v14.csv")))


def load_view(view):
    row = next(r for r in csv.DictReader(open(PROD / "output" / f"predictions_px_{view}.csv")) if r["ID"] == EXAMPLE)
    n = sum(1 for k in KEY if k["view"] == view)
    P = np.array([[float(row[f"x{k}"]), float(row[f"y{k}"])] for k in range(1, n + 1)])
    img = tifffile.imread(REPO / row["image"])
    return P, img, float(row["px_per_mm"])


def ids(view):
    return {k["id"]: int(k["number"]) for k in KEY if k["view"] == view}


# ---------------------------------------------------------------- geometry, as in lance_v14_morphometrics.R
def unit(a, b):
    u = b - a
    return u / np.linalg.norm(u)


def height_to(p, u, chain):
    """From point p along the normal to axis u, the nearest crossing of the polyline `chain` (R height_to)."""
    n = np.array([-u[1], u[0]]); best = None
    for q0, q1 in zip(chain[:-1], chain[1:]):
        e = q1 - q0
        den = n[0] * e[1] - n[1] * e[0]
        if abs(den) < 1e-12:
            continue
        s = ((q0 - p)[0] * e[1] - (q0 - p)[1] * e[0]) / den
        t = ((q0 - p)[0] * n[1] - (q0 - p)[1] * n[0]) / den
        if 0 <= t <= 1 and (best is None or abs(s) < abs(best)):
            best = s
    return None if best is None else p + best * n


def foot(p, a, b):
    """Foot of the perpendicular from p on the line a-b."""
    u = unit(a, b)
    return a + ((p - a) @ u) * u


def base(img, P, width=1800, pad=110, crop=None):
    f = Fig(img, P if crop is None else crop, width=width, pad=pad)
    f.points(P, color=(90, 90, 90), r=3)
    return f


def key_figure(view, P, img, path):
    f = Fig(img, P, width=1800, pad=110)
    for k in KEY:
        if k["view"] != view:
            continue
        q = P[int(k["number"]) - 1]
        if k["role"] == "semilandmark":
            f.points(q[None], color=(240, 240, 240), r=4, outline=(0, 0, 0))
            f.text(q, k["number"], (60, 60, 60), size=16, dx=6, dy=-18)
        else:
            f.points(q[None], color=(40, 110, 230) if k["role"] == "computed" else (220, 30, 30), r=6, outline=(0, 0, 0))
            f.text(q, f'{k["number"]} {k["id"]}', (0, 0, 0), size=20, dx=6, dy=-30)
    f.save(path)


# ================================================================================== lateral (Right, Left)
def lateral_figures(view):
    s = view[0]
    P, img, ppm = load_view(view)
    I = ids(view)
    q = lambda nm: P[I[nm] - 1]  # noqa: E731
    right = view == "Right"
    vdist = list(range(25, 29)) if right else [25, 26]
    dorsal = list(range(29, 39)) if right else list(range(27, 37))
    window = list(range(39, 43)) if right else list(range(37, 41))
    ch = lambda nums: P[[n - 1 for n in nums]]  # noqa: E731
    dorsal_chain = ch([18] + dorsal + [1])
    u = unit(q(f"{s}02"), q(f"{s}01"))
    tag = view.lower()
    key_figure(view, P, img, FIGDIR / f"lance_{tag}_fig1_key.jpg")

    # lengths and arcs
    f = base(img, P)
    f.line(q(f"{s}03"), q(f"{s}01"), C[0], "Face_length", label_t=0.55, label_dy=-34)
    f.line(q(f"{s}02"), q(f"{s}01"), C[1], "Blade_chord", dashed=True, label_t=0.45, label_dy=18)
    f.line(q(f"{s}03"), q(f"{s}02"), C[2], "Heel_length", label_dx=-150)
    f.poly(ch([3, 19, 20, 21, 22, 2]), C[3], "Heel_arc_length", label_at=0.8, label_dx=-60, label_dy=40)
    f.poly(ch([2, 23, 24, 4, 5, 6, 7, 8, 9, 10] + vdist + [1]), C[4], "Ventral_arc_length", label_at=0.55, label_dy=20)
    f.poly(dorsal_chain, C[5], "Dorsal_arc_length", label_at=0.3, label_dy=-36)
    f.line(q(f"{s}10"), q(f"{s}01"), C[6], "Tip_length", label_dy=20)
    f.line(q(f"{s}17"), q(f"{s}01"), C[7], "Window_to_apex", label_t=0.4, label_dy=-34)
    f.save(FIGDIR / f"lance_{tag}_fig2_lengths.jpg")

    # window: length, dorsal suture-end spacing, scallops
    f = base(img, P)
    f.line(q(f"{s}11"), q(f"{s}17"), C[0], "Cutwindow_length", label_t=0.6, label_dy=-44)
    for k in range(1, 7):
        f.line(P[10 + k - 1], P[11 + k - 1], C[3], "Dorsal_arc1_width" if k == 1 else None, label_dx=-280, label_dy=-10, width=3)
    for k in range(1, 5):
        a, b, m = P[11 + k - 1], P[12 + k - 1], P[window[k - 1] - 1]
        fo = foot(m, a, b)
        f.line(m, fo, C[4], "Scallop1_amp" if k == 1 else None, width=5, label_dx=-60, label_dy=40)
    f.save(FIGDIR / f"lance_{tag}_fig3a_window.jpg")

    # sutures: lengths, ventral spacing, angle (drawn for suture 4)
    f = base(img, P)
    for k in range(1, 7):
        f.line(P[10 + k - 1], P[3 + k - 1], C[1], "Sut1_length" if k == 1 else None, label_dx=-200, label_dy=-10)
    for k in range(1, 7 if right else 6):
        f.line(P[3 + k - 1], P[4 + k - 1], C[2], "Arc1_width" if k == 1 else None, label_dy=18)
    k = 4; v = P[3 + k - 1]; top = P[10 + k - 1]
    f.line(v, v + u * 160, C[5], None, dashed=True)
    f.angle(v, v + u * 160, top, C[5], "Sut4_angle", radius=60)
    f.save(FIGDIR / f"lance_{tag}_fig3b_sutures.jpg")

    # heights
    f = base(img, P)
    f.poly(dorsal_chain, (150, 150, 150), width=2)
    f.line(q(f"{s}11"), q(f"{s}18"), C[0], "Height_window_start", label_dx=-240)
    for k in range(3, 8 if right else 7):
        a = P[3 + k - 1]; h = height_to(a, u, dorsal_chain)
        if h is not None:
            f.line(a, h, C[1], f"Height_sut{k}" if k == 3 else None, label_t=0.5, label_dx=-160)
    if not right:
        a = P[10 - 1]; h = height_to(a, u, dorsal_chain)
        if h is not None:
            f.line(a, h, C[6], "Height_split", label_t=0.5, label_dx=10)
        f.line(P[9 - 1], P[10 - 1], C[7], "Sut6_to_split", label_dy=24)
    for k in range(2, 7):
        a = P[10 + k - 1]; h = height_to(a, u, dorsal_chain)
        if h is not None:
            f.line(a, h, C[2], f"Band_height_sut{k}" if k == 2 else None, width=5, label_dx=10)
    f.save(FIGDIR / f"lance_{tag}_fig4_heights.jpg")
    return ppm


# ============================================================================================== bottom
def bottom_figures():
    P, img, ppm = load_view("Bottom")
    b = lambda n: P[n - 1]  # noqa: E731
    key_figure("Bottom", P, img, FIGDIR / "lance_bottom_fig1_key.jpg")
    mid = (b(1) + b(2)) / 2
    u = unit(b(11), mid)

    tip = P[[n - 1 for n in [1, 2, 3, 4, 7, 8, 14] + list(range(15, 27))]]
    f = base(img, P, width=1600, pad=200, crop=tip)
    f.line(b(14), b(1), C[0], "Long_half_length", label_t=0.6, label_dy=-44)
    f.line(b(14), b(2), C[1], "Short_half_length", label_t=0.6, label_dy=22)
    f.line(b(1), b(2), C[2], "Apex_separation", label_t=0.5, label_dx=16, label_dy=-12)
    f.line(b(14), mid, C[3], "Fork_depth", dashed=True, label_t=0.0, label_dx=-200, label_dy=-12)
    f.angle(b(14), b(15), b(18), C[5], None, radius=90)
    f.text(b(14), "Tip_divergence_angle", C[5], dx=-120, dy=70)
    f.angle(b(14), b(1), b(2), C[4], "Inner_angle", radius=220)
    f.save(FIGDIR / "lance_bottom_fig2_fork.jpg")

    f = base(img, P)
    f.line(b(11), mid, (120, 120, 120), "lance axis B11 -> apex midpoint", dashed=True, label_t=0.62, label_dy=6)
    for n, col, lab, dy in ((1, C[0], "Dorsal_length_long", -36), (2, C[1], "Dorsal_length_short", 18)):
        fo = b(11) + ((b(n) - b(11)) @ u) * u
        f.line(b(11) + np.array([0, -55 if n == 1 else 55]), fo + np.array([0, -55 if n == 1 else 55]), col, lab, label_dy=dy)
        f.line(b(n), fo, col, None, dashed=True)
    f.line(b(12), b(13), C[2], "Shoulder_span", label_t=0.85, label_dx=-200, label_dy=-10)
    for n, col, lab in ((12, C[3], "Shoulder_lat_long"), (13, C[4], "Shoulder_lat_short")):
        fo = b(11) + ((b(n) - b(11)) @ u) * u
        f.line(b(n), fo, col, lab, width=5, label_t=0.3, label_dx=-250 if n == 12 else 14, label_dy=-10)
        f.line(b(11), fo, col, None, dashed=True)
    f.save(FIGDIR / "lance_bottom_fig3_axis.jpg")

    f = base(img, P, pad=260)
    for k in range(1, 5):
        f.line(b(2 + k), b(6 + k), C[0], "Width_sut1" if k == 1 else ("Width_sut4" if k == 4 else None),
               label_t=1.0, label_dx=-70, label_dy=14)
    f.poly(P[[n - 1 for n in [3, 21, 22, 23, 1]]], C[1], "Tip_outer_arc_long", label_at=0.6, label_dy=-44)
    f.poly(P[[n - 1 for n in [7, 24, 25, 26, 2]]], C[2], "Tip_outer_arc_short", label_dy=16)
    f.poly(P[[n - 1 for n in [14, 15, 16, 17, 1]]], C[3], "Tip_medial_arc_long", label_at=0.0, label_dx=-330, label_dy=-44)
    f.poly(P[[n - 1 for n in [14, 18, 19, 20, 2]]], C[4], "Tip_medial_arc_short", label_at=0.0, label_dx=-330, label_dy=8)
    f.poly(P[[n - 1 for n in [6, 27, 28, 29, 30, 31, 32, 12]]], C[5], "Side_arc_long", label_at=0.75, label_dx=-60, label_dy=-40)
    f.poly(P[[n - 1 for n in [10, 33, 34, 35, 36, 37, 38, 13]]], C[6], "Side_arc_short", label_at=0.75, label_dx=-60, label_dy=16)
    f.save(FIGDIR / "lance_bottom_fig4_widths.jpg")
    return ppm


ppm_r = lateral_figures("Right")
lateral_figures("Left")
bottom_figures()

# ================================================================================================ document
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Inches(1))
normal = doc.styles["Normal"]
normal.font.name = "Calibri"; normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri"); normal.font.size = Pt(11)
for lvl in (1, 2):
    st = doc.styles[f"Heading {lvl}"]
    st.font.name, st.font.size, st.font.bold = "Calibri", Pt(13 if lvl == 1 else 12), True
    st.font.color.rgb = RGBColor(0, 0, 0); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def para(text, bold=False, italic=False, size=None):
    q = doc.add_paragraph(); r_ = q.add_run(text); r_.bold, r_.italic = bold, italic
    if size:
        r_.font.size = Pt(size)
    return q


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hex_fill)
    tcPr.append(s)


def table(header, rows, widths=(1.8, 1.5, 3.2)):
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = h; shade(c, "D9D9D9"); c.paragraphs[0].runs[0].bold = True
    for rw in rows:
        cells = t.add_row().cells
        for i, v in enumerate(rw):
            cells[i].text = v
    for rr in t.rows:
        for i, w in enumerate(widths):
            rr.cells[i].width = Inches(w)
            for q in rr.cells[i].paragraphs:
                for run in q.runs:
                    run.font.size = Pt(9.5)
    doc.add_paragraph()


def picture(path, caption, width=6.5):
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(caption, italic=True, size=9.5)


H = ["Measure (column)", "Points (Right view)", "Definition"]

doc.add_heading("Lance Protocol v1.4: Linear Measurements", level=0)
para("Definitions of the linear measures written by production/lance_v14_morphometrics.R into "
     "PRIME_Lance<View>_Pheno_v14.csv (functions lateral_measures and bottom_measures). They are computed from the "
     "CNN configurations in Lance<View>_v14_XY.csv.")
for t_ in [
    "Units: millimetres, from each image's own scale bar. Angles are in degrees.",
    "Column names carry the view as a suffix: _right, _left, _bot (e.g. Face_length_right).",
    "Point numbers follow landmark_key_v14.csv and are given in brackets, e.g. R04 (4). The Left view uses the same numbers for L01–L18. Its curve numbers differ (see the Left section).",
    "Lance axis for the lateral views: heel-ventral junction → apex (R02 → R01). Heights are measured perpendicular to this axis, from a point up to where that perpendicular meets the dorsal curve (R18 → dorsal semilandmarks → R01).",
    "Distances are straight lines unless the definition says 'along' a curve (sum of the segments through the curve's points).",
    f"Figures: CNN landmarks of the N. pinetum parent {EXAMPLE}, drawn exactly as the script computes the measures.",
]:
    doc.add_paragraph(t_, style="List Bullet")

# ----- Right
doc.add_heading("Right view (long half)", level=1)
picture(FIGDIR / "lance_right_fig1_key.jpg", "Figure 1. Right view: the 42 points (red = anchors, blue = computed R18, white = semilandmarks).")
doc.add_heading("Lengths and curves", level=2)
picture(FIGDIR / "lance_right_fig2_lengths.jpg", "Figure 2. Right view: lengths and curve lengths.")
table(H, [
    ["Face_length", "R03 (3)–R01 (1)", "Proximal heel notch to apex: whole face length."],
    ["Blade_chord", "R02 (2)–R01 (1)", "Heel-ventral junction to apex (straight). Also the lance axis."],
    ["Heel_length", "R03 (3)–R02 (2)", "Heel notch to heel-ventral junction (straight)."],
    ["Heel_arc_length", "R03, 19–22, R02", "Heel notch to heel-ventral junction along the heel margin."],
    ["Ventral_arc_length", "R02, 23–24, R04–R10, 25–28, R01", "Blade length along the ventral (sclerotized) margin, heel junction to apex."],
    ["Dorsal_arc_length", "R18, 29–38, R01", "Blade length along the dorsal curve, from its start R18 to the apex."],
    ["Tip_length", "R10 (10)–R01 (1)", "Last ventral suture end (suture 7) to apex."],
    ["Window_to_apex", "R17 (17)–R01 (1)", "Distal end of the cuticular window to apex."],
])
doc.add_heading("Window and sutures", level=2)
picture(FIGDIR / "lance_right_fig3a_window.jpg", "Figure 3. Right view: the cuticular window: length, spacing of the dorsal suture ends, and scallop amplitudes.")
picture(FIGDIR / "lance_right_fig3b_sutures.jpg", "Figure 4. Right view: suture lengths, spacing of the ventral suture ends, and the suture angle (drawn for suture 4).")
table(H, [
    ["Cutwindow_length", "R11 (11)–R17 (17)", "Window, proximal to distal end."],
    ["Sutk_length (k = 1–6)", "R(10+k)–R(3+k)", "Suture k, dorsal end on the window margin to ventral end (suture 1 uses the window's proximal end R11)."],
    ["Sutk_angle (k = 1–6)", "R(3+k)→R(10+k) vs R02→R01", "Inclination of suture k to the lance axis, 0–90°."],
    ["Arck_width (k = 1–6)", "R(3+k)–R(4+k)", "Spacing of ventral suture ends k and k+1 (R04–R05 … R09–R10)."],
    ["Dorsal_arck_width (k = 1–6)", "R(10+k)–R(11+k)", "Spacing of the dorsal suture ends along the window (R11–R12 … R16–R17)."],
    ["Scallopk_amp (k = 1–4)", "window point 38+k vs chord R(11+k)–R(12+k)", "Height of the k-th window scallop crest above the straight line between its two suture ends."],
])
doc.add_heading("Heights", level=2)
picture(FIGDIR / "lance_right_fig4_heights.jpg", "Figure 5. Right view: heights, measured perpendicular to the lance axis up to the dorsal curve (grey).")
table(H, [
    ["Height_window_start", "R11 (11)–R18 (18)", "Window's proximal end to the dorsal edge (perpendicular to the axis by construction of R18)."],
    ["Height_sutk (k = 3–7)", "R(3+k) → dorsal curve", "Lance height at ventral suture end k. Sutures 1–2 lie proximal to R18, so they have no height."],
    ["Band_height_sutk (k = 2–6)", "R(10+k) → dorsal curve", "Dorsal suture end (window margin) to the dorsal edge: the 'cutting depth' band above the window."],
])

# ----- Left
doc.add_heading("Left view (short half)", level=1)
picture(FIGDIR / "lance_left_fig1_key.jpg", "Figure 6. Left view: the 40 points.")
para("Every Right-view measure is also computed for the Left view with the same definition, suffix _left, and L in place of R. The differences follow from the Left view's points:")
for t_ in [
    "L10 (10) is the ventral split point, where the short half separates, not a seventh suture. So there are only Arc1–Arc5 widths, and Height_sut runs k = 3–6.",
    "The curves are numbered differently: L.vdist = 25–26 (2 points), L.dorsal = 27–36, L.window = 37–40 (Scallopk_amp uses window point 36+k).",
    "Two extra measures (Figure 7): Sut6_to_split = L09 (9)–L10 (10), the last suture to the split point; and Height_split = L10 → dorsal curve, the lance height at the split point.",
]:
    doc.add_paragraph(t_, style="List Bullet")
picture(FIGDIR / "lance_left_fig4_heights.jpg", "Figure 7. Left view: heights, including the Left-only Height_split and Sut6_to_split.")

# ----- Bottom
doc.add_heading("Bottom view (both halves)", level=1)
picture(FIGDIR / "lance_bottom_fig1_key.jpg", "Figure 8. Bottom view: the 38 points (B14 fork crotch computed).")
HB = ["Measure (column)", "Points", "Definition"]
doc.add_heading("Tips and fork", level=2)
picture(FIGDIR / "lance_bottom_fig2_fork.jpg", "Figure 9. Bottom view: half lengths, apex separation, fork depth and angles.")
table(HB, [
    ["Long_half_length", "B14 (14)–B01 (1)", "Fork crotch to the long-half apex."],
    ["Short_half_length", "B14 (14)–B02 (2)", "Fork crotch to the short-half apex."],
    ["Apex_separation", "B01 (1)–B02 (2)", "Distance between the two apices."],
    ["Fork_depth", "B14 → midpoint of B01, B02", "Fork crotch to the midpoint between the apices."],
    ["Inner_angle", "B01–B14–B02", "Angle at the fork crotch between the two apices."],
    ["Tip_divergence_angle", "15–B14–18", "Angle at the crotch between the first medial semilandmark of each tip (how the tips diverge just past the fork)."],
])
doc.add_heading("Positions along the lance axis", level=2)
picture(FIGDIR / "lance_bottom_fig3_axis.jpg", "Figure 10. Bottom view: the lance axis (B11 → apex midpoint), positions along it and distances from it.")
table(HB, [
    ["Dorsal_length_long / _short", "B01 / B02 projected on the axis", "Length of the whole structure along the axis, from the basal notch B11 to each apex."],
    ["Apex_offset", "Dorsal_length_long − _short", "How far the long apex extends beyond the short one (half-length asymmetry)."],
    ["Shoulder_span", "B12 (12)–B13 (13)", "Distance between the two shoulders."],
    ["Shoulder_pos_long / _short", "B12 / B13 projected on the axis", "Position of each shoulder along the axis from B11."],
    ["Shoulder_to_apex_long / _short", "Dorsal_length − Shoulder_pos", "Shoulder to apex, along the axis."],
    ["Shoulder_lat_long / _short", "B12 / B13 to the axis", "Perpendicular distance of each shoulder from the axis."],
    ["Shoulder_asym_pos / _lat", "long − short", "Shoulder asymmetry along and across the axis."],
])
doc.add_heading("Widths and curves", level=2)
picture(FIGDIR / "lance_bottom_fig4_widths.jpg", "Figure 11. Bottom view: widths across the suture pairs and curve lengths.")
table(HB, [
    ["Width_sutk (k = 1–4)", "B(2+k)–B(6+k)", "Width across the k-th suture pair counted from the apex (B03–B07 … B06–B10)."],
    ["Shaft_taper", "(Shoulder_span − Width_sut1) / axis distance", "Width gained per mm from suture 1 to the shoulders (axis positions of the suture-1 midpoint and the shoulder midpoint)."],
    ["Tip_outer_arc_long / _short", "B03, 21–23, B01 / B07, 24–26, B02", "Outer margin of each tip, suture 1 to apex."],
    ["Tip_medial_arc_long / _short", "B14, 15–17, B01 / B14, 18–20, B02", "Medial (inner) margin of each tip, fork crotch to apex."],
    ["Side_arc_long / _short", "B06, 27–32, B12 / B10, 33–38, B13", "Outer margin from the 4th suture to the shoulder."],
])

doc.add_heading("Notes", level=1)
for t_ in [
    "Shape is analysed separately, by Procrustes superimposition of all points with the semilandmarks sliding. The linear measures are for univariate and QTL models.",
    "Every length scales with lance size. Use centroid_size_<view> (in the same table) as a covariate for size-corrected comparisons.",
    "A height is NA when the perpendicular from its point does not meet the dorsal curve (rare, e.g. in very bent specimens).",
]:
    doc.add_paragraph(t_, style="List Bullet")
para(f"Generated by genai_lance_landmarking/tools/docgen/make_linear_measures_docx.py. Example scale (Right): {ppm_r:.1f} px/mm.", italic=True, size=9)
zoom = doc.settings.element.find(qn("w:zoom"))  # python-docx's template omits the required w:percent
if zoom is not None and zoom.get(qn("w:percent")) is None:
    zoom.set(qn("w:percent"), "100")
doc.save(OUT)
print("wrote", OUT)
